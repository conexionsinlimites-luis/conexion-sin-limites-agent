# tests/eval_valentina.py — Suite de evaluacion de Valentina y el modo dueno
# Conexion Sin Limites

"""
Corre casos de prueba fijos, basados en bugs reales encontrados en
produccion (calculo de decos DirecTV, comuna sola, fechas de vencimiento
hardcodeadas, ruteo cruzado DirecTV/VTR/Movistar, router del modo dueno).

Dos categorias:

  DETERMINISTA — no llama a ningun modelo. Gratis, corre en milisegundos.
                 Prueba codigo Python puro (parseo de fechas, extraccion de
                 marcadores, ruteo por keywords) y el texto estatico del
                 catalogo. Los que tocan la BD insertan filas de prueba
                 (prefijo de telefono "999...") y las borran al terminar.

  MODELO       — llama al modelo real (Sonnet para el prompt de ventas,
                 Haiku para el router del modo dueno). Prueba si Valentina
                 "entiende" algo, no solo si el codigo funciona -- por eso
                 no se puede reemplazar por logica determinista. Gasta
                 credito real (acotado: 1 mensaje corto por caso). Solo
                 corre si se pasa --con-modelo.

Uso:
    python tests/eval_valentina.py                # solo deterministas
    python tests/eval_valentina.py --con-modelo    # todos (gasta credito real)

Seguridad: los casos deterministas que tocan BD se niegan a correr si
DATABASE_URL apunta al host de produccion conocido (maglev.proxy.rlwy.net).
Configura una DATABASE_URL distinta (o vacia/local) antes de correr esto.
"""

import asyncio
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import yaml

from agent.config import DATABASE_URL, ANTHROPIC_API_KEY

_RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_HOST_PRODUCCION = "maglev.proxy.rlwy.net"
_TEST_PHONE_PREFIX = "999"
_TEST_MARKER = "__EVAL_TEST__"


@dataclass
class Resultado:
    nombre: str
    categoria: str  # "DETERMINISTA" | "MODELO"
    ok: bool
    detalle: str = ""


RESULTADOS: list[Resultado] = []


async def _caso(nombre: str, categoria: str, fn):
    try:
        await fn()
        RESULTADOS.append(Resultado(nombre, categoria, True))
    except AssertionError as e:
        RESULTADOS.append(Resultado(nombre, categoria, False, str(e)))
    except Exception as e:
        RESULTADOS.append(Resultado(nombre, categoria, False, f"ERROR INESPERADO: {e!r}"))


def _verificar_bd_no_produccion():
    if _HOST_PRODUCCION in (DATABASE_URL or ""):
        raise RuntimeError(
            f"DATABASE_URL apunta al host de producción conocido ({_HOST_PRODUCCION}). "
            "Este caso inserta y borra filas de prueba — configura una "
            "DATABASE_URL distinta antes de correr esta suite."
        )


async def _insertar_lead_fixture(
    conn, telefono, created_at, estado="nuevo", ultima_interaccion=None, subproducto="DirecTV"
):
    await conn.execute(
        """
        INSERT INTO leads (telefono, nombre, identificador, estado, score, subproducto,
                            origen, created_at, ultima_interaccion)
        VALUES ($1, $2, $3, $4, $5, $6, 'eval_test', $7, $8)
        ON CONFLICT (telefono) DO UPDATE SET
            estado = EXCLUDED.estado,
            created_at = EXCLUDED.created_at,
            ultima_interaccion = EXCLUDED.ultima_interaccion,
            subproducto = EXCLUDED.subproducto
        """,
        telefono, _TEST_MARKER, f"EVAL-{telefono}", estado, 10, subproducto,
        created_at, ultima_interaccion or created_at,
    )


async def _limpiar_fixtures():
    from agent.database import get_pool
    pool = await get_pool()
    async with pool.acquire() as conn:
        await conn.execute("DELETE FROM historial_mensajes WHERE telefono LIKE $1", f"{_TEST_PHONE_PREFIX}%")
        await conn.execute(
            "DELETE FROM leads WHERE nombre = $1 OR telefono LIKE $2", _TEST_MARKER, f"{_TEST_PHONE_PREFIX}%"
        )


# ════════════════════════════════════════════════════════════════
# CATEGORÍA A — DETERMINISTAS
# ════════════════════════════════════════════════════════════════

async def caso_rango_fecha_chile():
    from agent import crm

    desde, hasta = crm.rango_fecha_chile("2026-09-07", "2026-09-10")
    dias = (hasta - desde).total_seconds() / 86400
    assert desde < hasta, f"desde ({desde}) debería ser antes que hasta ({hasta})"
    assert 3.9 <= dias <= 4.0, f"el rango 07→10 sept debería cubrir ~4 días, cubrió {dias:.2f}"

    # Bug real de esta noche: fecha_hasta ausente colapsaba el rango a un
    # solo día en vez de cubrir el día completo (00:00 a 23:59:59 Chile).
    desde2, hasta2 = crm.rango_fecha_chile("2026-09-07")
    assert (hasta2 - desde2).total_seconds() > 86000, (
        f"un rango de un solo día debería cubrir ~24h, cubrió "
        f"{(hasta2 - desde2).total_seconds() / 3600:.1f}h"
    )


async def caso_leads_nuevos_fixtures():
    _verificar_bd_no_produccion()
    from agent import crm
    from agent.database import get_pool

    await _limpiar_fixtures()
    try:
        pool = await get_pool()
        async with pool.acquire() as conn:
            await _insertar_lead_fixture(conn, "9990000001", datetime(2026, 9, 8, 15, 0, 0))
            await _insertar_lead_fixture(conn, "9990000002", datetime(2026, 9, 10, 10, 0, 0))
            await _insertar_lead_fixture(conn, "9990000003", datetime(2026, 9, 1, 10, 0, 0))  # fuera de rango

        resultado = await crm.leads_nuevos("2026-09-07", "2026-09-10")
        assert resultado["total"] == 2, f"esperaba 2 leads nuevos en el rango, obtuvo {resultado['total']}"
        telefonos = {d["telefono"] for d in resultado["detalle"]}
        assert telefonos == {"9990000001", "9990000002"}, f"detalle inesperado: {telefonos}"
    finally:
        await _limpiar_fixtures()


async def caso_contactos_activos_fixtures():
    _verificar_bd_no_produccion()
    from agent import crm
    from agent.database import get_pool

    await _limpiar_fixtures()
    try:
        pool = await get_pool()
        async with pool.acquire() as conn:
            # Lead viejo (created_at fuera de rango) que escribió DENTRO del período
            # -- exactamente el caso real: "quién me escribió hoy y ayer".
            await _insertar_lead_fixture(conn, "9990000004", datetime(2026, 1, 1, 10, 0, 0))
            await conn.execute(
                "INSERT INTO historial_mensajes (telefono, rol, mensaje, timestamp) "
                "VALUES ($1, 'user', 'hola', $2)",
                "9990000004", datetime(2026, 9, 9, 23, 30, 0),
            )
            # Mensaje fuera del período -- no debe contar
            await conn.execute(
                "INSERT INTO historial_mensajes (telefono, rol, mensaje, timestamp) "
                "VALUES ($1, 'user', 'hola vieja', $2)",
                "9990000005", datetime(2026, 8, 1, 10, 0, 0),
            )

        resultado = await crm.contactos_activos("2026-09-09", "2026-09-10")
        assert resultado["total"] == 1, f"esperaba 1 contacto activo, obtuvo {resultado['total']}"
        assert resultado["detalle"][0]["telefono"] == "9990000004", resultado["detalle"]
    finally:
        await _limpiar_fixtures()


async def caso_ventas_cerradas_detalle():
    _verificar_bd_no_produccion()
    from agent import crm
    from agent.database import get_pool

    await _limpiar_fixtures()
    try:
        pool = await get_pool()
        async with pool.acquire() as conn:
            await _insertar_lead_fixture(
                conn, "9990000006", datetime(2026, 9, 1, 10, 0, 0),
                estado="cerrado", ultima_interaccion=datetime(2026, 9, 8, 12, 0, 0),
                subproducto="VTR",
            )

        resultado = await crm.ventas_cerradas("2026-09-07", "2026-09-10")
        assert resultado["total"] == 1, f"esperaba 1 venta cerrada, obtuvo {resultado['total']}"
        assert resultado["detalle"][0]["telefono"] == "9990000006", resultado["detalle"]
        assert resultado["detalle"][0]["subproducto"] == "VTR", resultado["detalle"]
    finally:
        await _limpiar_fixtures()


async def caso_ruteo_cruzado_directv_vtr():
    from agent.main import _clasificar_producto_lead

    lead_dtv = {"lead_resumen": "Cliente interesado en DirecTV full", "subproducto": ""}
    lead_vtr = {"lead_resumen": "Preguntó por VTR y Movistar", "subproducto": ""}
    lead_otro = {"lead_resumen": "Preguntó por Entel", "subproducto": ""}
    lead_falso_positivo = {"lead_resumen": "quiere buen servicio y atención", "subproducto": ""}

    assert _clasificar_producto_lead(lead_dtv, []) == "directv"
    assert _clasificar_producto_lead(lead_vtr, []) == "vtr_movistar"
    assert _clasificar_producto_lead(lead_otro, []) == "otro"
    assert _clasificar_producto_lead(lead_falso_positivo, []) == "otro", (
        "falso positivo: 'servicio' no debería matchear 'vtr' por substring"
    )


async def caso_cambio_modo_producto():
    from agent.main import _detectar_cambio_modo_producto

    assert _detectar_cambio_modo_producto("quiero que respondas solo dtv") == "directv"
    assert _detectar_cambio_modo_producto("solo vtr y movistar por ahora") == "vtr_movistar"
    assert _detectar_cambio_modo_producto("vuelve a todos los productos") == "todos"
    assert _detectar_cambio_modo_producto("hola como estas") is None


async def caso_extraer_alerta_supervisor():
    from agent.main import _extraer_alerta

    texto = (
        "Perfecto, ya tengo todo!\n"
        "[ALERTA_SUPERVISOR|nombre=Juan Perez|tel=56912345678|dir=Calle Falsa 123, Providencia]"
    )
    limpio, datos = _extraer_alerta(texto)
    assert datos is not None, "debería haber detectado la alerta"
    assert datos["nombre"] == "Juan Perez", datos
    assert datos["tel"] == "56912345678", datos
    assert "ALERTA_SUPERVISOR" not in limpio, "el marcador no debe llegar al cliente"


async def caso_extraer_alerta_duda():
    from agent.main import _extraer_alerta_duda

    texto = "Buena pregunta, déjame confirmar eso con el equipo [ALERTA_DUDA|pregunta=cual es el precio de WOM]"
    limpio, pregunta = _extraer_alerta_duda(texto)
    assert pregunta == "cual es el precio de WOM", pregunta
    assert "ALERTA_DUDA" not in limpio, "el marcador no debe llegar al cliente"

    limpio2, pregunta2 = _extraer_alerta_duda("mensaje normal sin marcador")
    assert pregunta2 is None, "no debería detectar nada sin el marcador"


async def caso_validar_campos_obligatorios():
    from agent.main import _validar_campos_obligatorios

    incompleto = {"tipo": "venta", "nombre": "Juan", "campos_faltantes": []}
    faltantes = _validar_campos_obligatorios(incompleto)
    assert "telefono" in faltantes, faltantes
    assert "producto" in faltantes, faltantes       # es venta -> producto es obligatorio
    assert "rut" in faltantes, faltantes             # RUT siempre obligatorio en una venta
    assert "internet_o_tv" in faltantes, faltantes   # hace falta saber qué se vendió
    assert "forma_pago" in faltantes, faltantes      # forma_pago ahora es SIEMPRE obligatorio

    # DirecTV completo -> nada pendiente (no pide carnet_foto_recibida, ahí basta el RUT)
    directv_completo = {
        "tipo": "venta", "nombre": "Juan", "telefono": "56912345678", "rut": "12.345.678-9",
        "producto": "DirecTV", "incluye_internet": True, "incluye_tv": True,
        "forma_pago": "PAT/PAC", "campos_faltantes": [],
    }
    assert _validar_campos_obligatorios(directv_completo) == [], _validar_campos_obligatorios(directv_completo)

    # VTR sin carnet_foto_recibida: sí debe quedar pendiente
    vtr_sin_carnet = {
        "tipo": "venta", "nombre": "Juan", "telefono": "56912345678", "rut": "12.345.678-9",
        "producto": "VTR", "incluye_internet": True, "incluye_tv": False,
        "forma_pago": "efectivo", "campos_faltantes": [],
    }
    faltantes_vtr = _validar_campos_obligatorios(vtr_sin_carnet)
    assert "carnet_foto_recibida" in faltantes_vtr, faltantes_vtr

    # VTR completo, con carnet_foto_recibida=False EXPLÍCITO (no es lo mismo
    # que "no se sabe") -> nada pendiente
    vtr_completo = {**vtr_sin_carnet, "carnet_foto_recibida": False}
    assert _validar_campos_obligatorios(vtr_completo) == [], _validar_campos_obligatorios(vtr_completo)

    # DirecTV sin carnet_foto_recibida: NO debe pedirlo (solo VTR/Movistar)
    assert "carnet_foto_recibida" not in _validar_campos_obligatorios(directv_completo)


async def caso_carga_pendiente_mensaje_no_relacionado():
    """
    Regresión del bug real: con una carga pendiente, un mensaje que el
    extractor marca como mensaje_no_relacionado NO debe pisar los datos
    pendientes -- la respuesta debe avisar del pendiente Y atender el
    mensaje nuevo, nunca ignorarlo repitiendo solo la pregunta vieja.
    """
    import agent.main as main

    telefono = "eval-test-carga-no-relacionado"
    datos_previos = {
        "tipo": "venta", "nombre": "Pedro Ramirez", "telefono": "56911112222",
        "producto": "VTR", "incluye_internet": True, "incluye_tv": True,
        "forma_pago": "efectivo", "rut": "11.111.111-1",
        "campos_faltantes": ["carnet_foto_recibida"], "listo_para_confirmar": False,
    }
    main._CARGA_PENDIENTE[telefono] = {"datos": datos_previos, "listo": False}

    bloque = SimpleNamespace(
        type="tool_use", name="registrar_carga",
        input={**datos_previos, "mensaje_no_relacionado": True},
    )
    respuesta_extraccion_falsa = SimpleNamespace(content=[bloque])

    async def _fake_create(*args, **kwargs):
        return respuesta_extraccion_falsa

    mensajes_enviados = []

    async def _fake_enviar_mensaje(tel, mensaje):
        mensajes_enviados.append((tel, mensaje))
        return True

    original_create = main.claude_client.messages.create
    original_enviar = main.proveedor.enviar_mensaje
    main.claude_client.messages.create = _fake_create
    main.proveedor.enviar_mensaje = _fake_enviar_mensaje
    try:
        await main._procesar_mensaje_dueño(telefono, "gracias")
        assert telefono in main._CARGA_PENDIENTE, "la carga pendiente no debería haberse perdido"
        assert main._CARGA_PENDIENTE[telefono]["datos"] == datos_previos, "los datos pendientes no deberían cambiar"
    finally:
        main.claude_client.messages.create = original_create
        main.proveedor.enviar_mensaje = original_enviar
        main._CARGA_PENDIENTE.pop(telefono, None)

    assert len(mensajes_enviados) == 1, mensajes_enviados
    texto_enviado = mensajes_enviados[0][1]
    assert "Pedro Ramirez" in texto_enviado, texto_enviado
    assert "pendiente" in texto_enviado.lower(), texto_enviado
    assert "De nada" in texto_enviado, texto_enviado


async def caso_carga_pendiente_si_incompleto_no_ejecuta_de_una():
    """
    Regresión: con campos_faltantes todavía pendientes (ej.
    carnet_foto_recibida), un "sí" del dueño NO debe guardar la venta de
    inmediato -- debe completar el campo que faltaba y mostrar la ficha
    final de confirmación, esperando un "sí" SEPARADO para recién guardar.
    """
    import agent.main as main

    telefono = "eval-test-carga-si-incompleto"
    datos_previos = {
        "tipo": "venta", "nombre": "Maria Torres", "telefono": "56933334444",
        "producto": "Movistar", "incluye_internet": True, "incluye_tv": False,
        "forma_pago": "PAT/PAC", "rut": "22.222.222-2",
        "campos_faltantes": ["carnet_foto_recibida"], "listo_para_confirmar": False,
    }
    main._CARGA_PENDIENTE[telefono] = {"datos": datos_previos, "listo": False}

    datos_completados = {
        **datos_previos, "carnet_foto_recibida": True,
        "campos_faltantes": [], "listo_para_confirmar": True,
    }
    bloque = SimpleNamespace(type="tool_use", name="registrar_carga", input=datos_completados)
    respuesta_falsa = SimpleNamespace(content=[bloque])

    async def _fake_create(*args, **kwargs):
        return respuesta_falsa

    llamadas_registrar_venta = []

    async def _fake_registrar_venta(**kwargs):
        llamadas_registrar_venta.append(kwargs)

    mensajes_enviados = []

    async def _fake_enviar_mensaje(tel, mensaje):
        mensajes_enviados.append((tel, mensaje))
        return True

    original_create = main.claude_client.messages.create
    original_registrar = main.crm.registrar_venta
    original_enviar = main.proveedor.enviar_mensaje
    main.claude_client.messages.create = _fake_create
    main.crm.registrar_venta = _fake_registrar_venta
    main.proveedor.enviar_mensaje = _fake_enviar_mensaje
    try:
        await main._procesar_mensaje_dueño(telefono, "si ya la tengo")
        assert not llamadas_registrar_venta, "no debería haber guardado la venta todavía, solo completar el campo"
        assert main._CARGA_PENDIENTE[telefono]["listo"] is True, main._CARGA_PENDIENTE.get(telefono)
    finally:
        main.claude_client.messages.create = original_create
        main.crm.registrar_venta = original_registrar
        main.proveedor.enviar_mensaje = original_enviar
        main._CARGA_PENDIENTE.pop(telefono, None)

    assert len(mensajes_enviados) == 1, mensajes_enviados
    assert "confirmas" in mensajes_enviados[0][1].lower(), mensajes_enviados


async def caso_carga_pendiente_si_completo_ejecuta():
    """Cuando la ficha final ya está lista (listo=True), un 'sí' SÍ debe ejecutar/guardar."""
    import agent.main as main

    telefono = "eval-test-carga-si-completo"
    datos_listos = {
        "tipo": "venta", "nombre": "Juan Perez", "telefono": "56955556666",
        "producto": "DirecTV", "incluye_internet": True, "incluye_tv": True,
        "forma_pago": "PAT/PAC", "rut": "33.333.333-3",
        "campos_faltantes": [], "listo_para_confirmar": True,
    }
    main._CARGA_PENDIENTE[telefono] = {"datos": datos_listos, "listo": True}

    llamadas_ejecutar = []

    async def _fake_ejecutar_carga(datos):
        llamadas_ejecutar.append(datos)
        return "Listo, creé un lead nuevo para Juan Pérez ✅"

    mensajes_enviados = []

    async def _fake_enviar_mensaje(tel, mensaje):
        mensajes_enviados.append((tel, mensaje))
        return True

    original_ejecutar = main._ejecutar_carga_dueño
    original_enviar = main.proveedor.enviar_mensaje
    main._ejecutar_carga_dueño = _fake_ejecutar_carga
    main.proveedor.enviar_mensaje = _fake_enviar_mensaje
    try:
        await main._procesar_mensaje_dueño(telefono, "si")
    finally:
        main._ejecutar_carga_dueño = original_ejecutar
        main.proveedor.enviar_mensaje = original_enviar
        main._CARGA_PENDIENTE.pop(telefono, None)

    assert len(llamadas_ejecutar) == 1, llamadas_ejecutar
    assert telefono not in main._CARGA_PENDIENTE


async def caso_cancelacion_carga_frase_completa():
    """
    'cancela eso' y 'olvida esa carga' deben cancelar la carga pendiente --
    antes solo funcionaban las palabras sueltas exactas ('no', 'cancelar').
    """
    import agent.main as main

    for frase in ["cancela eso", "olvida esa carga", "cancelar por favor"]:
        telefono = "eval-test-cancelacion-frase"
        main._CARGA_PENDIENTE[telefono] = {"datos": {"nombre": "Test"}, "listo": False}

        mensajes_enviados = []

        async def _fake_enviar_mensaje(tel, mensaje):
            mensajes_enviados.append((tel, mensaje))
            return True

        original_enviar = main.proveedor.enviar_mensaje
        main.proveedor.enviar_mensaje = _fake_enviar_mensaje
        try:
            await main._procesar_mensaje_dueño(telefono, frase)
        finally:
            main.proveedor.enviar_mensaje = original_enviar
            main._CARGA_PENDIENTE.pop(telefono, None)

        assert mensajes_enviados == [(telefono, "Cancelado, no se guardó nada.")], (frase, mensajes_enviados)


async def caso_carga_con_vtr_solo_no_cambia_modo_producto():
    """
    Regresión de un bug real en producción: "carga venta: Pedro Ramírez
    56922222222 RUT 98765432-1 VTR, solo internet, efectivo" contiene
    "solo" + "VTR" -- coincidía con las keywords de cambio de modo de
    producto, y ese detector se evaluaba ANTES que el de carga, cambiando
    el modo global de Valentina a "solo VTR y Movistar" en vez de iniciar
    el flujo de carga (afectó a clientes reales preguntando por DirecTV).
    """
    import agent.main as main

    telefono = "eval-test-carga-vtr-solo"
    texto = "carga venta: Pedro Ramirez 56922222222 RUT 98765432-1 VTR, solo internet, efectivo"

    datos_extraidos = {
        "tipo": "venta", "nombre": "Pedro Ramirez", "telefono": "56922222222",
        "rut": "98765432-1", "producto": "VTR",
        "incluye_internet": True, "incluye_tv": False, "forma_pago": "efectivo",
        "campos_faltantes": ["carnet_foto_recibida"], "listo_para_confirmar": False,
    }
    bloque = SimpleNamespace(type="tool_use", name="registrar_carga", input=datos_extraidos)
    respuesta_falsa = SimpleNamespace(content=[bloque])

    async def _fake_create(*args, **kwargs):
        return respuesta_falsa

    llamadas_modo_producto = []

    async def _fake_actualizar_modo(modo, cliente_slug=None):
        llamadas_modo_producto.append(modo)

    mensajes_enviados = []

    async def _fake_enviar_mensaje(tel, mensaje):
        mensajes_enviados.append((tel, mensaje))
        return True

    original_create = main.claude_client.messages.create
    original_actualizar_modo = main.prompt_builder.actualizar_modo_producto
    original_enviar = main.proveedor.enviar_mensaje
    main.claude_client.messages.create = _fake_create
    main.prompt_builder.actualizar_modo_producto = _fake_actualizar_modo
    main.proveedor.enviar_mensaje = _fake_enviar_mensaje
    try:
        await main._procesar_mensaje_dueño(telefono, texto)
        assert telefono in main._CARGA_PENDIENTE, "debería haber entrado al flujo de carga"
    finally:
        main.claude_client.messages.create = original_create
        main.prompt_builder.actualizar_modo_producto = original_actualizar_modo
        main.proveedor.enviar_mensaje = original_enviar
        main._CARGA_PENDIENTE.pop(telefono, None)

    assert not llamadas_modo_producto, f"el modo de producto NO debía cambiar, pero se llamó con: {llamadas_modo_producto}"
    assert len(mensajes_enviados) == 1, mensajes_enviados


async def caso_router_dueno_multi_llamada():
    import agent.main as main

    bloque_hoy = SimpleNamespace(
        type="tool_use", name="contactos_activos",
        input={"fecha_desde": "2026-09-10", "fecha_hasta": "2026-09-10"},
    )
    bloque_ayer = SimpleNamespace(
        type="tool_use", name="contactos_activos",
        input={"fecha_desde": "2026-09-09", "fecha_hasta": "2026-09-09"},
    )
    respuesta_falsa = SimpleNamespace(content=[bloque_hoy, bloque_ayer])

    async def _fake_create(*args, **kwargs):
        return respuesta_falsa

    original_create = main.claude_client.messages.create
    main.claude_client.messages.create = _fake_create
    try:
        llamadas = await main._rutear_consulta_dueño("quien me escribio hoy y ayer", "eval-test-tel")
    finally:
        main.claude_client.messages.create = original_create

    assert len(llamadas) == 2, f"esperaba 2 llamadas (hoy y ayer), obtuvo {len(llamadas)}"
    nombres = [n for n, _ in llamadas]
    assert nombres == ["contactos_activos", "contactos_activos"], nombres


async def caso_historial_dueño_limite_turnos():
    """
    _guardar_turno_historial_dueño debe quedarse solo con los últimos N
    turnos (descartando los más viejos) y truncar respuestas largas, para
    que un resumen extenso (ej. export en chat) no infle el contexto de
    turnos futuros.
    """
    import agent.main as main

    telefono = "eval-test-historial-limite"
    main._HISTORIAL_DUEÑO.pop(telefono, None)
    try:
        for i in range(main._LIMITE_TURNOS_HISTORIAL + 3):
            main._guardar_turno_historial_dueño(telefono, f"pregunta {i}", f"respuesta {i}")
        turnos = main._HISTORIAL_DUEÑO[telefono]
        assert len(turnos) == main._LIMITE_TURNOS_HISTORIAL, turnos
        # deben quedar los ÚLTIMOS N, no los primeros
        assert turnos[0]["pregunta"] == "pregunta 3", turnos
        assert turnos[-1]["pregunta"] == f"pregunta {main._LIMITE_TURNOS_HISTORIAL + 2}", turnos

        respuesta_larga = "x" * 1000
        main._guardar_turno_historial_dueño(telefono, "otra pregunta", respuesta_larga)
        assert len(main._HISTORIAL_DUEÑO[telefono][-1]["respuesta"]) == main._LIMITE_CARACTERES_RESPUESTA_HISTORIAL
    finally:
        main._HISTORIAL_DUEÑO.pop(telefono, None)


async def caso_router_sin_historial_previo():
    """Sin turnos guardados, el router debe mandar SOLO la pregunta actual -- el comportamiento base no debe cambiar."""
    import agent.main as main

    telefono = "eval-test-sin-historial"
    main._HISTORIAL_DUEÑO.pop(telefono, None)

    llamada_capturada = {}
    bloque = SimpleNamespace(type="tool_use", name="dato_no_registrado", input={"razon": "n/a"})
    respuesta_falsa = SimpleNamespace(content=[bloque])

    async def _fake_create(*args, **kwargs):
        llamada_capturada.update(kwargs)
        return respuesta_falsa

    original_create = main.claude_client.messages.create
    main.claude_client.messages.create = _fake_create
    try:
        await main._rutear_consulta_dueño("pregunta nueva sin contexto previo", telefono)
    finally:
        main.claude_client.messages.create = original_create

    mensajes = llamada_capturada.get("messages", [])
    assert mensajes == [{"role": "user", "content": "pregunta nueva sin contexto previo"}], mensajes


async def caso_router_usa_historial_para_referencias():
    """
    Regresión del bug real: el router debe incluir los turnos previos de
    _HISTORIAL_DUEÑO como mensajes anteriores en la llamada a Claude --
    sin esto, no puede resolver referencias como "esos dos leads" a una
    pregunta anterior ("leads del 8 a la fecha").
    """
    import agent.main as main

    telefono = "eval-test-historial-router"
    main._HISTORIAL_DUEÑO.pop(telefono, None)
    main._guardar_turno_historial_dueño(
        telefono, "leads del 8 a la fecha",
        "2 leads entraron entre el 8 y el 11 de septiembre.",
    )

    llamada_capturada = {}
    bloque = SimpleNamespace(
        type="tool_use", name="leads_nuevos",
        input={"fecha_desde": "2026-09-08", "fecha_hasta": "2026-09-11"},
    )
    respuesta_falsa = SimpleNamespace(content=[bloque])

    async def _fake_create(*args, **kwargs):
        llamada_capturada.update(kwargs)
        return respuesta_falsa

    original_create = main.claude_client.messages.create
    main.claude_client.messages.create = _fake_create
    try:
        await main._rutear_consulta_dueño("los numeros de esos dos leads", telefono)
    finally:
        main.claude_client.messages.create = original_create
        main._HISTORIAL_DUEÑO.pop(telefono, None)

    mensajes = llamada_capturada.get("messages", [])
    assert len(mensajes) == 3, mensajes
    assert mensajes[0] == {"role": "user", "content": "leads del 8 a la fecha"}, mensajes
    assert mensajes[1] == {"role": "assistant", "content": "2 leads entraron entre el 8 y el 11 de septiembre."}, mensajes
    assert mensajes[2] == {"role": "user", "content": "los numeros de esos dos leads"}, mensajes


async def caso_responder_consulta_dueño_guarda_turno():
    """
    _responder_consulta_dueño debe guardar automáticamente el turno en
    _HISTORIAL_DUEÑO al terminar -- sin esto la memoria nunca se llena,
    aunque el router ya sepa leerla.
    """
    import agent.main as main

    telefono = "eval-test-guardar-turno"
    main._HISTORIAL_DUEÑO.pop(telefono, None)

    bloque = SimpleNamespace(type="tool_use", name="dato_no_registrado", input={"razon": "prueba"})
    respuesta_router_falsa = SimpleNamespace(content=[bloque])

    async def _fake_create(*args, **kwargs):
        return respuesta_router_falsa

    original_create = main.claude_client.messages.create
    main.claude_client.messages.create = _fake_create
    try:
        respuesta = await main._responder_consulta_dueño("una pregunta cualquiera", telefono)
    finally:
        main.claude_client.messages.create = original_create

    turnos = main._HISTORIAL_DUEÑO.get(telefono, [])
    assert len(turnos) == 1, turnos
    assert turnos[0]["pregunta"] == "una pregunta cualquiera", turnos
    assert turnos[0]["respuesta"] == respuesta, turnos

    main._HISTORIAL_DUEÑO.pop(telefono, None)


async def caso_cortesia_deteccion_exacta():
    """
    _es_cortesia debe reconocer SOLO el mensaje completo de cortesía (con
    tolerancia a mayúsculas, espacios y signos de puntuación), y NUNCA una
    pregunta real que solo contenga una de esas palabras como parte de una
    frase más larga (falso positivo real que hay que evitar).
    """
    from agent.main import _es_cortesia, _respuesta_cortesia

    for frase in ["gracias", "muchas gracias", "gracias!", "¡gracias!", "  gracias  ", "GRACIAS"]:
        assert _es_cortesia(frase.lower()), frase
        assert _respuesta_cortesia(frase.lower()) == "De nada 😊", frase

    for frase in ["ok", "listo", "dale", "perfecto", "listo!", "Dale."]:
        assert _es_cortesia(frase.lower()), frase
        assert _respuesta_cortesia(frase.lower()) == "👍", frase

    no_son_cortesia = [
        "ok pero cuantos leads entraron",
        "dale la direccion de Juan",
        "gracias por la venta de ayer, cuanto fue",
        "listo para cerrar esta venta",
    ]
    for frase in no_son_cortesia:
        assert not _es_cortesia(frase.lower()), frase


async def caso_cortesia_no_llega_al_router():
    """
    Regresión de integración: un "gracias" no debe disparar NINGUNA
    llamada al router de Haiku (el motor de consulta nunca debe
    ejecutarse) — debe responderse directo, sin gastar ni una llamada.
    """
    import agent.main as main

    llamadas_router = []

    async def _fake_create(*args, **kwargs):
        llamadas_router.append(kwargs)
        bloque = SimpleNamespace(
            type="tool_use", name="dato_no_registrado",
            input={"razon": "no debería haber llegado aquí"},
        )
        return SimpleNamespace(content=[bloque])

    mensajes_enviados = []

    async def _fake_enviar_mensaje(telefono, mensaje):
        mensajes_enviados.append((telefono, mensaje))
        return True

    original_create = main.claude_client.messages.create
    original_enviar = main.proveedor.enviar_mensaje
    main.claude_client.messages.create = _fake_create
    main.proveedor.enviar_mensaje = _fake_enviar_mensaje
    try:
        await main._procesar_mensaje_dueño("eval-test-cortesia", "gracias!")
    finally:
        main.claude_client.messages.create = original_create
        main.proveedor.enviar_mensaje = original_enviar

    assert not llamadas_router, f"el router se llamó pese a ser una cortesía: {llamadas_router}"
    assert mensajes_enviados == [("eval-test-cortesia", "De nada 😊")], mensajes_enviados


_PATRON_FECHA_DURA = re.compile(
    r"\b\d{1,2}\s+de\s+(enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
    r"septiembre|octubre|noviembre|diciembre)\b",
    re.IGNORECASE,
)


async def caso_sin_fechas_vencidas_hardcodeadas():
    archivos = ["config/prompts.yaml", "config/prompt_base.txt", "config/comportamiento.md"]
    encontrados = []
    for ruta in archivos:
        ruta_completa = os.path.join(_RAIZ, ruta)
        if not os.path.exists(ruta_completa):
            continue
        with open(ruta_completa, "r", encoding="utf-8") as f:
            texto = f.read()
        for m in _PATRON_FECHA_DURA.finditer(texto):
            inicio, fin = max(0, m.start() - 40), min(len(texto), m.end() + 40)
            encontrados.append(f"{ruta}: ...{texto[inicio:fin]}...")
    assert not encontrados, "fecha de vencimiento hardcodeada encontrada:\n" + "\n".join(encontrados)


async def caso_comision_directv_menos_de_14_puntos():
    from agent import crm

    assert crm.calcular_sueldo_directv(10) == 200_000, crm.calcular_sueldo_directv(10)
    assert crm.calcular_sueldo_directv(5.5) == 110_000, crm.calcular_sueldo_directv(5.5)


async def caso_comision_directv_exactamente_14_puntos():
    from agent import crm

    assert crm.calcular_sueldo_directv(14) == 350_000, crm.calcular_sueldo_directv(14)


async def caso_comision_directv_mas_de_14_puntos():
    from agent import crm

    assert crm.calcular_sueldo_directv(16) == 390_000, crm.calcular_sueldo_directv(16)
    assert crm.calcular_sueldo_directv(20) == 470_000, crm.calcular_sueldo_directv(20)


async def caso_puntos_venta_duo_con_pat_pac():
    from agent import crm

    duo_pat_pac = {"incluye_internet": True, "incluye_tv": True, "forma_pago": "PAT/PAC"}
    duo_sin_bono = {"incluye_internet": True, "incluye_tv": True, "forma_pago": "efectivo"}
    assert crm.puntos_venta_directv(duo_pat_pac) == 2.5, crm.puntos_venta_directv(duo_pat_pac)
    assert crm.puntos_venta_directv(duo_sin_bono) == 2.0, crm.puntos_venta_directv(duo_sin_bono)


async def caso_puntos_venta_solo_con_pat_pac():
    from agent import crm

    solo_internet_pat_pac = {"incluye_internet": True, "incluye_tv": False, "forma_pago": "pat/pac"}
    solo_tv_pat_pac = {"incluye_internet": False, "incluye_tv": True, "forma_pago": "PAT/PAC"}
    solo_sin_bono = {"incluye_internet": True, "incluye_tv": False, "forma_pago": "efectivo"}
    assert crm.puntos_venta_directv(solo_internet_pat_pac) == 2.0, crm.puntos_venta_directv(solo_internet_pat_pac)
    assert crm.puntos_venta_directv(solo_tv_pat_pac) == 2.0, crm.puntos_venta_directv(solo_tv_pat_pac)
    assert crm.puntos_venta_directv(solo_sin_bono) == 1.5, crm.puntos_venta_directv(solo_sin_bono)

    # Venta sin internet ni TV -- no cubierta por la regla, debe fallar fuerte
    try:
        crm.puntos_venta_directv({"incluye_internet": False, "incluye_tv": False, "forma_pago": "efectivo"})
        raise AssertionError("debería haber lanzado ValueError para venta sin internet ni TV")
    except ValueError:
        pass


async def caso_bono_rally_califica():
    """6+ ventas Dúo y >=30% del total DirecTV en PAT/PAC -> +$100.000, sumado al sueldo por puntos."""
    from agent import crm

    ventas = [
        {"incluye_internet": True, "incluye_tv": True, "forma_pago": "PAT/PAC"}
        for _ in range(6)
    ]
    rally = crm.calcular_bono_rally_directv(ventas)
    assert rally["ventas_duo"] == 6, rally
    assert rally["porcentaje_pat_pac"] == 100.0, rally
    assert rally["califica"] is True, rally
    assert rally["bono"] == 100_000, rally

    comision = crm.calcular_comision_directv(ventas)
    assert comision["sueldo_total"] == comision["sueldo_por_puntos"] + 100_000, comision


async def caso_bono_rally_no_califica_por_ventas_duo():
    """5 ventas Dúo (todas PAT/PAC) -- no llega a las 6 requeridas, no califica aunque el % esté sobrado."""
    from agent import crm

    ventas = [
        {"incluye_internet": True, "incluye_tv": True, "forma_pago": "PAT/PAC"}
        for _ in range(5)
    ]
    rally = crm.calcular_bono_rally_directv(ventas)
    assert rally["ventas_duo"] == 5, rally
    assert rally["califica"] is False, rally
    assert rally["bono"] == 0, rally


async def caso_bono_rally_no_califica_por_porcentaje():
    """6 ventas Dúo PAT/PAC + 15 solo efectivo -- 6/21 = 28.6% < 30%, no califica."""
    from agent import crm

    ventas = (
        [{"incluye_internet": True, "incluye_tv": True, "forma_pago": "PAT/PAC"} for _ in range(6)]
        + [{"incluye_internet": True, "incluye_tv": False, "forma_pago": "efectivo"} for _ in range(15)]
    )
    rally = crm.calcular_bono_rally_directv(ventas)
    assert rally["ventas_duo"] == 6, rally
    assert rally["total_ventas"] == 21, rally
    assert rally["porcentaje_pat_pac"] < 30.0, rally
    assert rally["califica"] is False, rally
    assert rally["bono"] == 0, rally


async def caso_bono_rally_boundary_30_porciento():
    """6 ventas Dúo PAT/PAC + 14 solo efectivo -- 6/20 = exactamente 30%, sí califica (umbral inclusivo)."""
    from agent import crm

    ventas = (
        [{"incluye_internet": True, "incluye_tv": True, "forma_pago": "PAT/PAC"} for _ in range(6)]
        + [{"incluye_internet": True, "incluye_tv": False, "forma_pago": "efectivo"} for _ in range(14)]
    )
    rally = crm.calcular_bono_rally_directv(ventas)
    assert rally["porcentaje_pat_pac"] == 30.0, rally
    assert rally["califica"] is True, rally
    assert rally["bono"] == 100_000, rally


def _venta_duo(compania):
    return {"compania": compania, "incluye_internet": True, "incluye_tv": True}


def _venta_solo(compania):
    return {"compania": compania, "incluye_internet": True, "incluye_tv": False}


async def caso_comision_vtr_claro_tramo1():
    """
    25 RGU, todo dúo (tramo 1: 1-30). Nota: dúo aporta 2 RGU por venta, así
    que 25 exacto no es alcanzable solo con dúo -- uso 24 RGU (12 ventas),
    el total par más cercano, sin salir del tramo 1 (1-30).
    """
    from agent import crm

    ventas = [_venta_duo("VTR") for _ in range(12)]  # 12 x 2 RGU = 24 RGU
    resultado = crm.calcular_comision_vtr_claro(ventas)

    assert resultado["rgu_total"] == 24, resultado
    assert resultado["tarifa_duo"] == 70_000, resultado
    assert resultado["total_final"] == 12 * 70_000, resultado
    # En tramo 1, el pago inicial YA es la tarifa final -- no debería haber diferencia el día 5
    assert resultado["diferencia_dia_5"] == 0, resultado


async def caso_comision_vtr_claro_cruza_tramo2():
    """
    35 RGU (15 dúo = 30 RGU + 5 solo = 5 RGU) -> cruza a tramo 2 (31-50).
    Verifica que la diferencia del día 5 se aplique a TODAS las ventas del
    mes -- incluidas las 15 dúo Y las 5 solo -- no solo a una parte.
    """
    from agent import crm

    ventas = [_venta_duo("VTR") for _ in range(15)] + [_venta_solo("Claro") for _ in range(5)]
    resultado = crm.calcular_comision_vtr_claro(ventas)

    assert resultado["rgu_total"] == 35, resultado
    assert resultado["tarifa_duo"] == 75_000, resultado
    assert resultado["tarifa_solo_internet"] == 45_000, resultado

    total_final_esperado = 15 * 75_000 + 5 * 45_000
    pago_inicial_esperado = 15 * 70_000 + 5 * 40_000  # tramo 1 (base) para las 20 ventas
    assert resultado["total_final"] == total_final_esperado, resultado
    assert resultado["pago_inicial"] == pago_inicial_esperado, resultado
    assert resultado["diferencia_dia_5"] == total_final_esperado - pago_inicial_esperado, resultado


async def caso_comision_vtr_claro_tramo3():
    """55 RGU (25 dúo = 50 RGU + 5 solo = 5 RGU) -> tramo 3 (51+)."""
    from agent import crm

    ventas = [_venta_duo("VTR") for _ in range(25)] + [_venta_solo("VTR") for _ in range(5)]
    resultado = crm.calcular_comision_vtr_claro(ventas)

    assert resultado["rgu_total"] == 55, resultado
    assert resultado["tarifa_duo"] == 80_000, resultado
    assert resultado["tarifa_solo_internet"] == 50_000, resultado
    assert resultado["total_final"] == 25 * 80_000 + 5 * 50_000, resultado


async def caso_comision_movistar_tramo2():
    """55 RGU de Movistar (27 dúo = 54 RGU + 1 solo = 1 RGU) -> tramo 2 (51+)."""
    from agent import crm

    ventas = [_venta_duo("Movistar") for _ in range(27)] + [_venta_solo("Movistar")]
    resultado = crm.calcular_comision_movistar(ventas)

    assert resultado["rgu_total"] == 55, resultado
    assert resultado["tarifa_duo"] == 65_000, resultado
    assert resultado["tarifa_solo_internet"] == 45_000, resultado
    assert resultado["total_final"] == 27 * 65_000 + 1 * 45_000, resultado


async def caso_comision_vtr_claro_ventas_mezcladas():
    """
    Ventas de VTR y Claro en el mismo mes deben sumar al mismo contador
    combinado. Además confirma que una venta de Movistar metida por error
    en el grupo VTR/Claro sea rechazada (no se mezcla en silencio).
    """
    from agent import crm

    ventas = [_venta_duo("VTR") for _ in range(10)] + [_venta_solo("Claro") for _ in range(5)]
    resultado = crm.calcular_comision_vtr_claro(ventas)

    assert resultado["rgu_total"] == 25, resultado  # 10*2 + 5*1
    assert resultado["cantidad_ventas"] == 15, resultado
    assert resultado["total_final"] == 10 * 70_000 + 5 * 40_000, resultado

    try:
        crm.calcular_comision_vtr_claro(ventas + [_venta_duo("Movistar")])
        raise AssertionError("debería haber rechazado una venta de Movistar en el grupo VTR/Claro")
    except ValueError:
        pass


async def caso_comision_vtr_claro_desglose_por_compania():
    """
    Regresión de bug real: una venta VTR se reportó como "1 venta de
    Claro" porque el resultado agregado no traía ningún desglose por
    compañía -- el formateador tenía que adivinar. Ahora debe venir
    explícito en ventas_por_compania.
    """
    from agent import crm

    resultado_una = crm.calcular_comision_vtr_claro([_venta_duo("VTR")])
    assert resultado_una["ventas_por_compania"] == {"VTR": 1}, resultado_una

    ventas_mixtas = [_venta_duo("VTR"), _venta_duo("VTR"), _venta_solo("Claro")]
    resultado_mixto = crm.calcular_comision_vtr_claro(ventas_mixtas)
    assert resultado_mixto["ventas_por_compania"] == {"VTR": 2, "Claro": 1}, resultado_mixto

    resultado_movistar = crm.calcular_comision_movistar([_venta_duo("Movistar")])
    assert resultado_movistar["ventas_por_compania"] == {"Movistar": 1}, resultado_movistar


async def caso_generar_xlsx_es_archivo_real():
    """
    Regresión de bug real: WhatsApp/Meta Cloud API rechaza text/csv como
    tipo de documento (no está en su lista oficial de MIME types
    soportados). El archivo generado debe ser un XLSX real y abrible, no
    un CSV con la extensión cambiada.
    """
    import io as io_module
    from openpyxl import load_workbook
    from agent.main import _generar_xlsx

    filas = [
        {"nombre": "Juan Perez", "telefono": "56911112222", "compania": "VTR", "monto_venta": 38910},
        {"nombre": "Maria Torres", "telefono": "56933334444", "compania": "Movistar", "monto_venta": 25000},
    ]
    contenido = _generar_xlsx(filas)

    # Firma ZIP (todo XLSX real es un ZIP) -- un CSV con extensión cambiada
    # jamás empezaría con estos bytes.
    assert contenido[:2] == b"PK", "el archivo generado no tiene la firma ZIP de un XLSX real"

    wb = load_workbook(io_module.BytesIO(contenido))
    filas_leidas = list(wb.active.iter_rows(values_only=True))
    assert filas_leidas[0] == ("nombre", "telefono", "compania", "monto_venta"), filas_leidas[0]
    assert filas_leidas[1] == ("Juan Perez", "56911112222", "VTR", 38910), filas_leidas[1]
    assert filas_leidas[2] == ("Maria Torres", "56933334444", "Movistar", 25000), filas_leidas[2]


async def caso_texto_resumen_export_legible():
    """
    Regresión de bug real: el respaldo de texto (cuando falla el envío
    del archivo, o cuando el dueño pide "chat") mostraba columnas crudas
    de la BD (id, lead_id, created_at con timestamp completo). Ahora debe
    mostrar solo campos legibles para un humano.
    """
    from agent.main import _texto_resumen_export

    filas_ventas = [{
        "id": 42, "lead_id": 7, "telefono": "56911112222", "nombre": "Juan Perez",
        "compania": "VTR", "monto_venta": 38910,
        "created_at": "2026-09-12 15:00:00.123456", "actualizado_en": "2026-09-12 15:00:00.123456",
    }]
    texto = _texto_resumen_export(filas_ventas, "ventas")

    assert "Nombre: Juan Perez" in texto, texto
    assert "Teléfono: 56911112222" in texto, texto
    assert "Compañía: VTR" in texto, texto
    assert "$38.910" in texto, texto
    for campo_tecnico in ("id:", "lead_id:", "created_at:", "actualizado_en:"):
        assert campo_tecnico not in texto, f"'{campo_tecnico}' no debería aparecer: {texto}"

    filas_leads = [{
        "id": 1, "telefono": "56922223333", "nombre": "Pedro Soto",
        "subproducto": "DirecTV", "estado": "nuevo", "created_at": "2026-09-12 10:00:00",
    }]
    texto_leads = _texto_resumen_export(filas_leads, "leads")
    assert "Compañía: DirecTV" in texto_leads, texto_leads
    assert "Estado: nuevo" in texto_leads, texto_leads
    assert "created_at" not in texto_leads, texto_leads


CASOS_DETERMINISTAS = [
    ("rango_fecha_chile límites correctos", caso_rango_fecha_chile),
    ("leads_nuevos con fixtures", caso_leads_nuevos_fixtures),
    ("contactos_activos con fixtures", caso_contactos_activos_fixtures),
    ("ventas_cerradas trae detalle", caso_ventas_cerradas_detalle),
    ("ruteo cruzado DirecTV/VTR/Movistar", caso_ruteo_cruzado_directv_vtr),
    ("cambio de modo de producto", caso_cambio_modo_producto),
    ("extracción [ALERTA_SUPERVISOR]", caso_extraer_alerta_supervisor),
    ("extracción [ALERTA_DUDA]", caso_extraer_alerta_duda),
    ("validación de campos obligatorios (carga)", caso_validar_campos_obligatorios),
    ("carga pendiente: mensaje no relacionado no la pisa (regresión)", caso_carga_pendiente_mensaje_no_relacionado),
    ("carga pendiente: 'sí' incompleto no ejecuta de una (regresión)", caso_carga_pendiente_si_incompleto_no_ejecuta_de_una),
    ("carga pendiente: 'sí' completo sí ejecuta", caso_carga_pendiente_si_completo_ejecuta),
    ("cancelación de carga por frase completa", caso_cancelacion_carga_frase_completa),
    ("carga con 'VTR, solo internet' no cambia el modo de producto (regresión)", caso_carga_con_vtr_solo_no_cambia_modo_producto),
    ("router dueño ejecuta TODAS las llamadas (mock)", caso_router_dueno_multi_llamada),
    ("historial dueño respeta el límite de turnos", caso_historial_dueño_limite_turnos),
    ("router sin historial manda solo la pregunta actual", caso_router_sin_historial_previo),
    ("router usa historial para resolver referencias (regresión)", caso_router_usa_historial_para_referencias),
    ("_responder_consulta_dueño guarda el turno en el historial", caso_responder_consulta_dueño_guarda_turno),
    ("cortesías: detección exacta, sin falsos positivos", caso_cortesia_deteccion_exacta),
    ("cortesías: 'gracias' no llega al router (regresión)", caso_cortesia_no_llega_al_router),
    ("sin fechas de vencimiento hardcodeadas", caso_sin_fechas_vencidas_hardcodeadas),
    ("comisión DirecTV: menos de 14 puntos (proporcional)", caso_comision_directv_menos_de_14_puntos),
    ("comisión DirecTV: exactamente 14 puntos ($350.000)", caso_comision_directv_exactamente_14_puntos),
    ("comisión DirecTV: más de 14 puntos (extra por punto)", caso_comision_directv_mas_de_14_puntos),
    ("comisión DirecTV: puntos venta Dúo + PAT/PAC", caso_puntos_venta_duo_con_pat_pac),
    ("comisión DirecTV: puntos venta Solo + PAT/PAC", caso_puntos_venta_solo_con_pat_pac),
    ("bono Rally DirecTV: califica (6 dúo, 100% PAT/PAC)", caso_bono_rally_califica),
    ("bono Rally DirecTV: no califica por ventas dúo (5 < 6)", caso_bono_rally_no_califica_por_ventas_duo),
    ("bono Rally DirecTV: no califica por % PAT/PAC (28.6% < 30%)", caso_bono_rally_no_califica_por_porcentaje),
    ("bono Rally DirecTV: umbral exacto 30% (inclusivo)", caso_bono_rally_boundary_30_porciento),
    ("comisión VTR/Claro: tramo 1 (24 RGU, todo dúo)", caso_comision_vtr_claro_tramo1),
    ("comisión VTR/Claro: cruza a tramo 2, diferencia retroactiva", caso_comision_vtr_claro_cruza_tramo2),
    ("comisión VTR/Claro: tramo 3 (55 RGU)", caso_comision_vtr_claro_tramo3),
    ("comisión Movistar: tramo 2 (55 RGU)", caso_comision_movistar_tramo2),
    ("comisión VTR/Claro: ventas mezcladas VTR+Claro", caso_comision_vtr_claro_ventas_mezcladas),
    ("comisión VTR/Claro/Movistar: desglose por compañía (regresión)", caso_comision_vtr_claro_desglose_por_compania),
    ("exportar: xlsx generado es un archivo real (regresión)", caso_generar_xlsx_es_archivo_real),
    ("exportar: resumen de texto legible, sin columnas crudas (regresión)", caso_texto_resumen_export_legible),
]


# ════════════════════════════════════════════════════════════════
# CATEGORÍA B — REQUIEREN MODELO REAL (gastan crédito, solo con --con-modelo)
# ════════════════════════════════════════════════════════════════

def _cargar_system_prompt_real() -> str:
    with open(os.path.join(_RAIZ, "config", "prompts.yaml"), "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return data["system_prompt"]


async def _preguntar_sonnet(mensaje: str, max_tokens: int = 500) -> str:
    from anthropic import AsyncAnthropic

    client = AsyncAnthropic(api_key=ANTHROPIC_API_KEY)
    respuesta = await client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=max_tokens,
        system=_cargar_system_prompt_real(),
        messages=[{"role": "user", "content": mensaje}],
    )
    return "".join(b.text for b in respuesta.content if b.type == "text")


async def caso_precio_2_decos_adicionales():
    texto = await _preguntar_sonnet(
        "Hola, quiero cotizar el Duo DTH Esencial 800 Megas pagando en efectivo, "
        "y necesito agregar 2 decos adicionales. Cuanto seria el total mensual?"
    )
    assert "38.910" in texto, f"no calculó el total correcto. Respuesta:\n{texto}"
    assert "35.970" not in texto, f"repitió el error de aplicar descuento a ambos decos. Respuesta:\n{texto}"
    assert "3.980" not in texto, f"repitió el error de aplicar descuento a ambos decos. Respuesta:\n{texto}"


async def caso_comuna_sola_pide_calle_numero():
    texto = await _preguntar_sonnet("Hola, vivo en Providencia, cuanto me saldria el internet?")
    texto_low = texto.lower()
    assert "calle" in texto_low or "dirección" in texto_low or "direccion" in texto_low, (
        f"no pidió la dirección completa. Respuesta:\n{texto}"
    )
    assert "[alerta_supervisor" not in texto_low, (
        f"disparó la alerta de dirección completa con solo la comuna. Respuesta:\n{texto}"
    )


async def caso_wom_sin_precio_inventado():
    texto = await _preguntar_sonnet("Hola, cuanto cuesta el plan de WOM de 600 megas?")
    assert not re.search(r"\$\s?\d", texto), f"inventó un precio de WOM. Respuesta:\n{texto}"


async def caso_router_dueno_preguntas_reales():
    import agent.main as main

    casos = [
        ("cuantos leads entraron esta semana", "leads_nuevos"),
        ("cuantas ventas llevamos este mes", "ventas_cerradas"),
        ("cuanto pesa la luna", "dato_no_registrado"),
        ("busca el telefono de Juan Perez", "buscar_lead_por_nombre"),
    ]
    fallos = []
    for pregunta, esperado in casos:
        llamadas = await main._rutear_consulta_dueño(pregunta, "eval-test-tel")
        nombres = [n for n, _ in llamadas]
        if esperado not in nombres:
            fallos.append(f"'{pregunta}' -> esperaba incluir {esperado}, obtuvo {nombres}")
    assert not fallos, "fallos de ruteo:\n" + "\n".join(fallos)

    llamadas = await main._rutear_consulta_dueño("quien me escribio hoy y ayer", "eval-test-tel")
    assert len(llamadas) >= 2, f"esperaba al menos 2 llamadas para 'hoy y ayer', obtuvo {len(llamadas)}"


async def caso_catalogo_extensor_gratis_940():
    import agent.main as main

    pregunta = "cuanto por internet de 940 megas directv y un extensor wifi"
    llamadas = await main._rutear_consulta_dueño(pregunta, "eval-test-tel")
    nombres = [n for n, _ in llamadas]
    assert "consultar_catalogo" in nombres, f"esperaba consultar_catalogo, obtuvo {nombres}"

    resultados = []
    for nombre_funcion, parametros in llamadas:
        datos = await main._ejecutar_consulta_dueño(nombre_funcion, parametros)
        resultados.append((nombre_funcion, datos))
    respuesta = await main._formatear_respuesta_dueño(pregunta, resultados)

    assert "16.990" in respuesta, (
        f"no calculó el total correcto (debería ser $16.990 — el extensor es "
        f"gratis en plan 940). Respuesta:\n{respuesta}"
    )
    texto_low = respuesta.lower()
    assert "gratis" in texto_low or "$0" in respuesta or "sin costo" in texto_low, (
        f"no mencionó que el extensor es gratis en plan 940. Respuesta:\n{respuesta}"
    )


CASOS_MODELO = [
    ("2 decos adicionales DirecTV", caso_precio_2_decos_adicionales),
    ("comuna sola pide calle+número", caso_comuna_sola_pide_calle_numero),
    ("WOM sin precio inventado", caso_wom_sin_precio_inventado),
    ("router dueño con preguntas reales", caso_router_dueno_preguntas_reales),
    ("catálogo: extensor gratis en 940 DirecTV", caso_catalogo_extensor_gratis_940),
]


async def main():
    con_modelo = "--con-modelo" in sys.argv

    print("=== EVAL VALENTINA ===\n")

    for nombre, fn in CASOS_DETERMINISTAS:
        await _caso(nombre, "DETERMINISTA", fn)

    if con_modelo:
        print("(--con-modelo activo: los siguientes casos llaman a la API real y gastan crédito)\n")
        for nombre, fn in CASOS_MODELO:
            await _caso(nombre, "MODELO", fn)
    else:
        print("(saltando casos de MODELO — correr con --con-modelo para incluirlos)\n")

    for r in RESULTADOS:
        estado = "PASS" if r.ok else "FAIL"
        etiqueta = f"[{r.categoria}]"
        print(f"{etiqueta:15} {r.nombre:.<55} {estado}")
        if not r.ok:
            for linea in r.detalle.splitlines():
                print(f"    -> {linea}")

    total = len(RESULTADOS)
    pasaron = sum(1 for r in RESULTADOS if r.ok)
    print("-" * 70)
    if pasaron == total:
        print(f"{pasaron}/{total} pasaron")
    else:
        print(f"{pasaron}/{total} pasaron ({total - pasaron} fallaron — detalle arriba)")

    from agent.database import close_pool
    await close_pool()

    if pasaron < total:
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(main())
